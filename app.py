import os
from flask import Flask, render_template, request, jsonify
import openpyxl

app = Flask(__name__)

EXCEL_FILE = 'citas.xlsx'

PALETA_COLORES = ['#3788d8', '#28a745', '#dc3545', '#fd7e14', '#6f42c1', '#17a2b8', '#e83e8c', '#6c757d']

def asegurar_columna_estado():
    """Añade la columna Estado al Excel si no existe y marca las citas antiguas como Activo."""
    if not os.path.exists(EXCEL_FILE):
        inicializar_excel()
        return

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Citas"] if "Citas" in wb.sheetnames else wb.active

    encabezado_estado = None

    for col in range(1, ws.max_column + 1):
        if str(ws.cell(1, col).value or "").strip().lower() == "estado":
            encabezado_estado = col
            break

    if encabezado_estado is None:
        encabezado_estado = ws.max_column + 1
        ws.cell(1, encabezado_estado).value = "Estado"

    for fila in range(2, ws.max_row + 1):
        if ws.cell(fila, 1).value is not None and not ws.cell(fila, encabezado_estado).value:
            ws.cell(fila, encabezado_estado).value = "Activo"

    wb.save(EXCEL_FILE)


def obtener_columna_estado(ws):
    """Devuelve el número de columna donde está Estado."""
    for col in range(1, ws.max_column + 1):
        if str(ws.cell(1, col).value or "").strip().lower() == "estado":
            return col
    return None


def inicializar_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws_citas = wb.active
        ws_citas.title = "Citas"
        headers_citas = [
            "ID", "Nombre", "Apellidos", "Teléfono", "DNI", "Email",
            "Valor Total", "Monto Pagado", "Nº Sesiones Totales",
            "Nº Box", "Tatuador", "Fecha Inicio", "Observacion"
        ]
        ws_citas.append(headers_citas)
        
        ws_tat = wb.create_sheet(title="Tatuadores")
        ws_tat.append(["Nombre Tatuador"])
        wb.save(EXCEL_FILE)

def obtener_lista_tatuadores():
    if not os.path.exists(EXCEL_FILE):
        inicializar_excel()
        return ["Tatuador 1", "Tatuador 2"]
        
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    tatuadores = []
    
    # 1. Buscar en cualquier pestaña que contenga "tat" o "hoja" o en la segunda pestaña
    target_sheet = None
    for name in wb.sheetnames:
        if "tat" in name.lower() or "hoja" in name.lower() or "sheet" in name.lower():
            if name.lower() != "citas":
                target_sheet = wb[name]
                break
                
    if not target_sheet and len(wb.sheetnames) > 1:
        target_sheet = wb[wb.sheetnames[1]]
        
    if target_sheet:
        for row in target_sheet.iter_rows(min_row=1, values_only=True):
            if row and row[0] is not None:
                valor = str(row[0]).strip()
                if valor.lower() in ["nombre tatuador", "nombre", "tatuador", "tatuadores", "id"]:
                    continue
                if valor and valor not in tatuadores:
                    tatuadores.append(valor)

    # 2. Si no encontró en pestañas secundarias, buscar valores únicos en la columna de citas
    if not tatuadores and "Citas" in wb.sheetnames:
        ws = wb["Citas"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) > 10 and row[10] is not None:
                valor = str(row[10]).strip()
                if valor and valor not in tatuadores:
                    tatuadores.append(valor)

    print(">>> TATUADORES DETECTADOS EN EXCEL:", tatuadores)
    return tatuadores if tatuadores else ["Tatuador 1", "Tatuador 2"]

inicializar_excel()

@app.route('/')
def index():
    tatuadores = obtener_lista_tatuadores()
    return render_template('index.html', tatuadores=tatuadores)

@app.route('/api/citas', methods=['GET'])
def get_citas():
    if not os.path.exists(EXCEL_FILE):
        inicializar_excel()
        
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb["Citas"] if "Citas" in wb.sheetnames else wb.active
    
    tatuadores = obtener_lista_tatuadores()
    mapa_colores = {tat: PALETA_COLORES[i % len(PALETA_COLORES)] for i, tat in enumerate(tatuadores)}
    
    conteo_sesiones = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        tel = str(row[3]).strip() if row[3] else ""
        if tel:
            conteo_sesiones[tel] = conteo_sesiones.get(tel, 0) + 1

    citas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        
        valor_total = float(row[6] or 0)
        monto_pagado = float(row[7] or 0)
        sesiones_totales = int(row[8] or 1)
        tel = str(row[3]).strip() if row[3] else ""
        sesiones_agendadas = conteo_sesiones.get(tel, 1)
        sesiones_restantes = max(0, sesiones_totales - sesiones_agendadas)
        
        if monto_pagado >= valor_total and valor_total > 0:
            estado_pago = "Pagado"
        elif monto_pagado > 0:
            estado_pago = "Parcial"
        else:
            estado_pago = "Pendiente"
            
        tatuador = str(row[10]).strip() if row[10] else "Sin Asignar"
        color = mapa_colores.get(tatuador, '#3788d8')

        citas.append({
            'id': row[0],
            'title': f"{row[1]} {row[2]} (Box {row[9]})",
            'start': str(row[11]),
            'color': color,
            'extendedProps': {
                'id': row[0],
                'nombre': row[1],
                'apellidos': row[2],
                'telefono': row[3],
                'dni': row[4],
                'email': row[5],
                'valor_total': valor_total,
                'monto_pagado': monto_pagado,
                'sesiones_totales': sesiones_totales,
                'sesiones_agendadas': sesiones_agendadas,
                'sesiones_restantes': sesiones_restantes,
                'n_box': row[9],
                'tatuador': tatuador,
                'estado_pago': estado_pago,
                'observacion': row[12]
            }
        })
    return jsonify(citas)

@app.route('/api/guardar_cita', methods=['POST'])
def guardar_cita():
    data = request.json
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Citas"] if "Citas" in wb.sheetnames else wb.active
    
    nuevo_id = ws.max_row
    ws.append([
        nuevo_id,
        data.get('nombre'),
        data.get('apellidos'),
        data.get('telefono'),
        data.get('dni'),
        data.get('email'),
        float(data.get('valor_total', 0)),
        float(data.get('monto_pagado', 0)),
        int(data.get('n_sesiones_totales', 1)),
        data.get('n_box'),
        data.get('tatuador'),
        data.get('fecha_inicio'),
        data.get('observacion')
    ])
    wb.save(EXCEL_FILE)
    return jsonify({'success': True})

@app.route('/api/actualizar_pago', methods=['POST'])
def actualizar_pago():
    data = request.json
    cita_id = int(data.get('id'))
    nuevo_monto = float(data.get('monto_pagado', 0))
    
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Citas"] if "Citas" in wb.sheetnames else wb.active
    
    for row in ws.iter_rows(min_row=2):
        if row[0].value == cita_id:
            row[7].value = nuevo_monto
            break
            
    wb.save(EXCEL_FILE)
    return jsonify({'success': True})

@app.route('/api/finalizar_cita', methods=['POST'])
def finalizar_cita():
    data = request.json
    cita_id = int(data.get('id'))

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Citas"] if "Citas" in wb.sheetnames else wb.active

    columna_estado = obtener_columna_estado(ws)

    if columna_estado is None:
        columna_estado = ws.max_column + 1
        ws.cell(1, columna_estado).value = "Estado"

    encontrada = False

    for fila in range(2, ws.max_row + 1):
        if ws.cell(fila, 1).value == cita_id:
            ws.cell(fila, columna_estado).value = "Finalizado"
            encontrada = True
            break

    wb.save(EXCEL_FILE)

    return jsonify({
        'success': encontrada,
        'message': 'Cita finalizada correctamente' if encontrada else 'Cita no encontrada'
    })


@app.route('/api/eliminar_cita', methods=['POST'])
def eliminar_cita():
    data = request.json
    cita_id = int(data.get('id'))

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Citas"] if "Citas" in wb.sheetnames else wb.active

    encontrada = False

    for fila in range(2, ws.max_row + 1):
        if ws.cell(fila, 1).value == cita_id:
            ws.delete_rows(fila, 1)
            encontrada = True
            break

    wb.save(EXCEL_FILE)

    return jsonify({
        'success': encontrada,
        'message': 'Cita eliminada correctamente' if encontrada else 'Cita no encontrada'
    })


@app.route('/api/informe_mensual', methods=['GET'])
def informe_mensual():
    mes_filtro = request.args.get('mes')
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb["Citas"] if "Citas" in wb.sheetnames else wb.active
    
    tatuadores = obtener_lista_tatuadores()
    total_facturado = 0
    total_cobrado = 0
    por_tatuador = {t: {'facturado': 0, 'cobrado': 0, 'pendiente': 0} for t in tatuadores}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[11]:
            continue
        fecha_str = str(row[11])
        if fecha_str.startswith(mes_filtro):
            val = float(row[6] or 0)
            pag = float(row[7] or 0)
            tat = str(row[10]).strip() if row[10] else "Sin Asignar"

            total_facturado += val
            total_cobrado += pag

            if tat not in por_tatuador:
                por_tatuador[tat] = {'facturado': 0, 'cobrado': 0, 'pendiente': 0}

            por_tatuador[tat]['facturado'] += val
            por_tatuador[tat]['cobrado'] += pag
            por_tatuador[tat]['pendiente'] += (val - pag)

    return jsonify({
        'facturado': total_facturado,
        'cobrado': total_cobrado,
        'pendiente': total_facturado - total_cobrado,
        'por_tatuador': por_tatuador
    })

if __name__ == '__main__':
    app.run(debug=True)
